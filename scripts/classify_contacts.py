#!/usr/bin/env python3
"""Classify exported contacts without modifying the source list.

Input: CSV or XLSX contact export.
Output: CSV audit copy with exclusion reasons, identity tier, and resolved name.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any, Iterable

HERE = Path(__file__).resolve().parent
DEFAULT_CONFIG = HERE.parent / "references" / "keyword-taxonomy.json"
NAME_CANDIDATES = ["微信备注名", "姓名", "name", "contact_name", "display_name", "备注名"]
USERNAME_CANDIDATES = ["微信号", "username", "user_id", "wxid", "account", "handle"]
PHONE_CANDIDATES = ["手机号", "电话", "phone", "mobile", "联系方式"]
NOTES_CANDIDATES = ["备注", "notes", "note", "description", "标签"]


def normalize(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    return re.sub(r"\s+", " ", text).strip()


def read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise SystemExit("openpyxl is required for XLSX input") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [normalize(x) or f"column_{i + 1}" for i, x in enumerate(rows[0])]
        return headers, [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        headers = [normalize(x) for x in (reader.fieldnames or [])]
        return headers, [dict(row) for row in reader]


def pick(headers: Iterable[str], explicit: str | None, candidates: list[str]) -> str | None:
    headers = list(headers)
    if explicit:
        return explicit if explicit in headers else None
    lowered = {h.lower(): h for h in headers}
    for name in candidates:
        if name in headers:
            return name
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


def first_value(row: dict[str, Any], columns: list[str]) -> str:
    for col in columns:
        value = normalize(row.get(col, ""))
        if value:
            return value
    return ""


def phones_in(text: str, pattern: str) -> list[str]:
    return sorted(set(re.findall(pattern, text)))


def person_name_candidate(name: str, brand_terms: list[str], title_terms: list[str]) -> str:
    text = normalize(name)
    for prefix in ("AAAA", "AAA", "AA", "A"):
        if text.startswith(prefix) and len(text) > len(prefix):
            text = text[len(prefix):].strip()
    if not text:
        return ""
    if any(term in text for term in brand_terms):
        return ""
    if re.fullmatch(r"[\u4e00-\u9fa5]{2,4}", text):
        return text
    match = re.match(r"([\u4e00-\u9fa5]{1,3})(?:" + "|".join(map(re.escape, title_terms)) + ")", text)
    return match.group(0) if match else ""


def classify(text: str, taxonomy: dict[str, Any]) -> tuple[list[str], list[str], bool]:
    categories = taxonomy.get("exclude_categories", {})
    keep_terms = taxonomy.get("keep_terms", [])
    protected = [term for term in keep_terms if term in text]
    excluded: list[tuple[str, str, bool]] = []
    for key, rule in categories.items():
        hits = [kw for kw in rule.get("keywords", []) if kw in text]
        if hits:
            excluded.append((key, "/".join(hits), bool(rule.get("strong"))))
    if not excluded:
        return [], [], False
    if protected:
        unoverridable = [item for item in excluded if item[2] and not categories.get(item[0], {}).get("allow_keep_override", False)]
        if not unoverridable:
            return [], protected, False
    return [item[0] for item in excluded], [item[1] for item in excluded], False


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input")
    parser.add_argument("output")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG))
    parser.add_argument("--name-col")
    parser.add_argument("--username-col")
    parser.add_argument("--phone-col")
    parser.add_argument("--notes-cols", default="")
    parser.add_argument("--keep-core", action="store_true", default=True)
    parser.add_argument("--no-keep-core", dest="keep_core", action="store_false")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser()
    output_path = Path(args.output).expanduser()
    taxonomy = json.loads(Path(args.config).expanduser().read_text(encoding="utf-8"))
    headers, rows = read_rows(input_path)
    if not headers:
        raise SystemExit("input has no columns")
    name_col = pick(headers, args.name_col, NAME_CANDIDATES)
    username_col = pick(headers, args.username_col, USERNAME_CANDIDATES)
    phone_col = pick(headers, args.phone_col, PHONE_CANDIDATES)
    notes_cols = [c.strip() for c in args.notes_cols.split(",") if c.strip()]
    if not notes_cols:
        notes_cols = [pick(headers, None, NOTES_CANDIDATES) or ""]
        notes_cols = [c for c in notes_cols if c]
    if not name_col and not username_col:
        raise SystemExit("could not identify name or username columns")

    phone_regex = taxonomy["identity"]["phone_regex"]
    brand_terms = taxonomy["identity"].get("brand_terms", [])
    title_terms = taxonomy["identity"].get("title_terms", [])
    out_rows = []
    counts = {"excluded": 0, "tier1": 0, "tier2": 0, "tier3": 0}
    for row in rows:
        original_name = normalize(row.get(name_col, "")) if name_col else ""
        username = normalize(row.get(username_col, "")) if username_col else ""
        extra = " ".join(normalize(row.get(c, "")) for c in notes_cols)
        phone_source = normalize(row.get(phone_col, "")) if phone_col else ""
        phone_text = " ".join([original_name, username, phone_source, extra])
        phones = phones_in(phone_text, phone_regex)
        name = person_name_candidate(original_name, brand_terms, title_terms)
        exclude_categories, exclude_reasons, _ = classify(" ".join([original_name, username, extra]), taxonomy)
        if not args.keep_core:
            exclude_categories = list(exclude_categories)
            exclude_reasons = list(exclude_reasons)
        if exclude_categories:
            counts["excluded"] += 1
            name = ""
        tier = 1 if (name and phones) else 2 if name else 3
        counts[f"tier{tier}"] += 1
        out_rows.append({
            "original_name": original_name,
            "resolved_name": name,
            "username": username,
            "identity_tier": tier,
            "contact_info": ",".join(phones),
            "exclude": "true" if exclude_categories else "false",
            "exclude_categories": ",".join(exclude_categories),
            "exclude_reasons": ",".join(exclude_reasons),
            "protected_keep_terms": ",".join([t for t in taxonomy.get("keep_terms", []) if t in phone_text]),
            "notes": extra,
        })
    fieldnames = list(out_rows[0].keys()) if out_rows else [
        "original_name", "resolved_name", "username", "identity_tier", "contact_info",
        "exclude", "exclude_categories", "exclude_reasons", "protected_keep_terms", "notes",
    ]
    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(out_rows)
    print(json.dumps({"output": str(output_path), "rows": len(out_rows), "counts": counts}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())